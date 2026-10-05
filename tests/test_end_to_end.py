import re

from openpyxl import load_workbook

from trustdesk.evaluate import load_gold, retrieval_recall, run
from trustdesk.llm import Citation, DraftAnswer, FakeLLM
from trustdesk.questionnaire import OUT_COLUMNS, answers_for_export, read_questions, write_completed


def oracle_llm(gold):
    """Answers like a perfect model would: quotes the right source, abstains when it should."""
    by_q = {g["question"]: g for g in gold}

    def respond(question, sources):
        g = by_q[question]
        blocks = re.findall(r"\[(S\d+)\] \(\w[\w ]*: ([^ ]+) ›[^\n]*\)\n(.*?)(?=\n\n\[S\d+\]|\Z)", sources, re.S)
        for sid, doc, text in blocks:
            if g["expected_doc"] and doc == g["expected_doc"]:
                sentence = re.split(r"(?<=\.)\s", text.strip())[0]
                return DraftAnswer(status="answered", answer=g["expected"], explanation="We comply.",
                                   citations=[Citation(source=sid, quote=sentence)])
        return DraftAnswer(status="insufficient_evidence", explanation="Not covered.")

    return FakeLLM(respond)


def test_questionnaire_parsing(sample):
    qs = read_questions(sample / "questionnaires" / "acme_bank_2026.xlsx")
    assert len(qs) == 40 and qs[0].qid == "AB-001" and qs[0].row == 4


def test_pipeline_with_oracle_model(engine, sample):
    gold = load_gold(sample / "eval" / "questions.jsonl")
    engine.llm = oracle_llm(gold)
    s = run(engine, gold).summary()
    # Unanswerable questions never get answered, and no answer is high-confidence wrong.
    assert s["bluff_rate"] == 0.0 and s["high_confidence_wrong"] == 0
    # Accuracy is capped only by retrieval: the oracle answers whenever the right doc was retrieved.
    assert s["answer_accuracy"] >= 0.7


def test_export_adds_columns_and_preserves_original(engine, sample, tmp_path):
    gold = load_gold(sample / "eval" / "questions.jsonl")
    engine.llm = oracle_llm(gold)
    src = sample / "questionnaires" / "acme_bank_2026.xlsx"
    for q in read_questions(src)[:5]:
        checked, _ = engine.ask(q.text)
        engine.store.save_draft(src.name, q.qid, q.text, checked.to_dict())
    rows = engine.store.answers(src.name)
    engine.store.approve(src.name, rows[0]["qid"], "Yes", "Edited by reviewer.")
    out = write_completed(src, tmp_path / "out", answers_for_export(engine.store.answers(src.name)))
    ws = load_workbook(out).active
    header = [c.value for c in ws[3]]
    assert header[:5] == ["Ref", "Domain", "Question", "Vendor Response", "Comments"]
    assert header[5:] == OUT_COLUMNS
    assert ws.cell(row=4, column=7).value == "Edited by reviewer."
    assert ws.cell(row=4, column=10).value == "Approved"


def test_approved_answers_join_the_library(engine):
    before = engine.store.count()["library"]
    engine.add_to_library("Do you offer IP allowlisting?", "Yes", "Admins can restrict by IP.", "approved:x.xlsx")
    assert engine.store.count()["library"] == before + 1


def test_retrieval_recall_reported(engine, sample):
    r = retrieval_recall(engine, load_gold(sample / "eval" / "questions.jsonl"))
    assert set(r) == {"recall@1", "recall@3", "recall@5"} and r["recall@5"] >= r["recall@1"]
