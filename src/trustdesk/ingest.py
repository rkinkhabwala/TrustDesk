"""Turn documents into section-aware chunks.

Policies: .md / .txt / .pdf / .docx → chunks that remember their document and heading path, so a
citation can say "03-encryption-policy.md › Customer-managed keys".
Past questionnaires: .csv / .xlsx with Question + Answer columns → one library chunk per row.
"""
from __future__ import annotations

import csv
import hashlib
import re
from dataclasses import dataclass
from pathlib import Path

from .config import settings

POLICY_SUFFIXES = {".md", ".markdown", ".txt", ".pdf", ".docx"}


@dataclass
class Chunk:
    id: str
    doc: str
    section: str
    text: str
    kind: str  # "policy" | "library"

    @property
    def label(self) -> str:
        return f"{self.doc} › {self.section}" if self.section else self.doc


def _cid(doc: str, section: str, text: str) -> str:
    return hashlib.sha1(f"{doc}|{section}|{text}".encode()).hexdigest()[:16]


def _split_long(text: str, max_chars: int) -> list[str]:
    if len(text) <= max_chars:
        return [text]
    parts, cur = [], ""
    for para in re.split(r"\n\s*\n", text):
        if cur and len(cur) + len(para) + 2 > max_chars:
            parts.append(cur)
            cur = para
        else:
            cur = f"{cur}\n\n{para}" if cur else para
    if cur:
        parts.append(cur)
    out: list[str] = []
    for p in parts:  # a single giant paragraph: hard-split on sentences
        if len(p) <= max_chars:
            out.append(p)
            continue
        buf = ""
        for sent in re.split(r"(?<=[.!?])\s+", p):
            if buf and len(buf) + len(sent) + 1 > max_chars:
                out.append(buf)
                buf = sent
            else:
                buf = f"{buf} {sent}".strip()
        if buf:
            out.append(buf)
    return out


def chunk_markdown(doc: str, text: str, max_chars: int | None = None) -> list[Chunk]:
    """Split on headings; each chunk carries its heading path (e.g. 'Encryption Policy > Data at rest')."""
    max_chars = max_chars or settings.chunk_chars
    path: list[str] = []
    sections: list[tuple[str, list[str]]] = [("", [])]
    for line in text.splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            level, title = len(m.group(1)), m.group(2).strip()
            path = path[: level - 1] + [title]
            sections.append((" > ".join(path), []))
        else:
            sections[-1][1].append(line)
    chunks = []
    for section, lines in sections:
        body = "\n".join(lines).strip()
        if not body:
            continue
        for piece in _split_long(body, max_chars):
            chunks.append(Chunk(_cid(doc, section, piece), doc, section, piece, "policy"))
    return chunks


def _pdf_text(path: Path) -> list[tuple[str, str]]:
    from pypdf import PdfReader

    return [(f"page {i}", page.extract_text() or "") for i, page in enumerate(PdfReader(str(path)).pages, 1)]


def _docx_markdown(path: Path) -> str:
    import docx

    lines = []
    for p in docx.Document(str(path)).paragraphs:
        style = (p.style.name or "").lower()
        if style.startswith("heading") and style[-1:].isdigit():
            lines.append("#" * int(style[-1]) + " " + p.text)
        elif p.text.strip():
            lines.append(p.text)
            lines.append("")
    return "\n".join(lines)


def load_policy(path: Path) -> list[Chunk]:
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown", ".txt"}:
        return chunk_markdown(path.name, path.read_text(encoding="utf-8"))
    if suffix == ".docx":
        return chunk_markdown(path.name, _docx_markdown(path))
    if suffix == ".pdf":
        out = []
        for section, text in _pdf_text(path):
            for piece in _split_long(text.strip(), settings.chunk_chars):
                if piece:
                    out.append(Chunk(_cid(path.name, section, piece), path.name, section, piece, "policy"))
        return out
    raise ValueError(f"unsupported policy file: {path.name}")


def _rows(path: Path) -> list[dict]:
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            return list(csv.DictReader(f))
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    header_idx = next(i for i, r in enumerate(rows) if any(str(c or "").lower().startswith("question") for c in r))
    header = [str(c or "").strip() for c in rows[header_idx]]
    return [dict(zip(header, r)) for r in rows[header_idx + 1:]]


def load_library(path: Path) -> list[Chunk]:
    """Past questionnaire → one chunk per answered question."""
    chunks = []
    for i, row in enumerate(_rows(path), 1):
        q = next((str(v) for k, v in row.items() if k and "question" in k.lower() and v), "")
        a = next((str(v) for k, v in row.items() if k and ("answer" in k.lower() or "response" in k.lower()) and v), "")
        if q.strip() and a.strip():
            text = f"Q: {q.strip()}\nA: {a.strip()}"
            chunks.append(Chunk(_cid(path.name, str(i), text), path.name, f"row {i}", text, "library"))
    return chunks


def collect(policy_dir: Path | None, library_paths: list[Path]) -> list[Chunk]:
    chunks: list[Chunk] = []
    if policy_dir:
        for p in sorted(policy_dir.rglob("*")):
            if p.is_file() and p.suffix.lower() in POLICY_SUFFIXES:
                chunks += load_policy(p)
    for p in library_paths:
        chunks += load_library(p)
    return chunks
